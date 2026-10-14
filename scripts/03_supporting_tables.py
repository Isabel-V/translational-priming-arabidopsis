"""Supporting Information tables (Tables S1-S9) and Table 1 for the translational-priming Letter.

Requires the outputs of scripts 01, 02 (Fig. 2b source data), 04 (motif analysis) and the g:Profiler results in
data/gprofiler/results (script 05).
Output: supporting_information/Supporting_Information_Tables_S1-S9.xlsx
        results/Table1_protein_evidence.tsv
"""
import numpy as np
import pandas as pd
from scipy.stats import fisher_exact, hypergeom
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from te_common import master_table, SI, RESULTS, RAW, ROOT, TE_FILE, MOTIF_FILE, TE_HIGH

SI.mkdir(exist_ok=True); RESULTS.mkdir(exist_ok=True)
m = master_table()
m["mean_log2_te"] = m[["log2_te_root", "log2_te_shoot"]].mean(axis=1)
split = m.primed_any & (((m.te_class_root == "high-TE") & ~m.primed_root) | ((m.te_class_shoot == "high-TE") & ~m.primed_shoot))
m["priming_status"] = np.select([m.primed_any & ~split, split, m.high_te_any], ["primed", "primed (canonical high-TE in other tissue)", "canonical high-TE"], "")
m["primed_tissues"] = np.select([m.primed_both, m.primed_root, m.primed_shoot], ["root + shoot", "root", "shoot"], "")
m["high_confidence_primed"] = (m.primed_root & (m.ribo_reads_root >= 10)) | (m.primed_shoot & (m.ribo_reads_shoot >= 10))

cols = {
    "gene_model": "Gene model", "agi": "AGI locus", "description": "Description (Araport11)",
    "functional_category": "Functional category", "organelle": "Genome", "predicted_protein_length_aa": "Predicted protein length (aa)",
    "rna_reads_root": "RNA-seq reads, root", "ribo_reads_root": "Ribo-seq reads, root",
    "rna_rpkm_root": "RNA-seq RPKM, root", "ribo_rpkm_root": "Ribo-seq RPKM, root",
    "log2_rna_root": "log2 RNA-seq RPKM, root", "log2_te_root": "log2 TE, root", "te_class_root": "TE class, root",
    "low_te_zero_ribo_root": "No Ribo-seq reads (TE = 0), root",
    "rna_reads_shoot": "RNA-seq reads, shoot", "ribo_reads_shoot": "Ribo-seq reads, shoot",
    "rna_rpkm_shoot": "RNA-seq RPKM, shoot", "ribo_rpkm_shoot": "Ribo-seq RPKM, shoot",
    "log2_rna_shoot": "log2 RNA-seq RPKM, shoot", "log2_te_shoot": "log2 TE, shoot", "te_class_shoot": "TE class, shoot",
    "low_te_zero_ribo_shoot": "No Ribo-seq reads (TE = 0), shoot",
    "primed_root": "Primed, root", "primed_shoot": "Primed, shoot", "primed_tissues": "Primed in",
    "priming_status": "High-TE status", "high_confidence_primed": "Primed with >=10 Ribo-seq reads",
    "seq_set": "Sequence-analysis set (log2 RNA >= 2 both tissues)",
    "protein_root": "Protein detected, root", "protein_shoot": "Protein detected, shoot", "protein_any": "Protein detected (root or shoot)",
}
def tidy(df, extra=()):
    d = df[list(cols) + list(extra)].rename(columns=cols)
    for c in d.columns:
        if d[c].dtype == float:
            d[c] = d[c].round(3)
    return d

# ---- Table 1 -----------------------------------------------------------
det = (m.rna_detected_root | m.rna_detected_shoot) & (m.ribo_detected_root | m.ribo_detected_shoot)
canon = m.high_te_any & ~m.primed_any
def row(label, crit, mask):
    s = m[mask]; k = int(s.protein_any.sum())
    return dict(gene_set=label, criterion=crit, n=len(s), with_protein=k, pct=round(100 * k / len(s), 1), odds_ratio="", p_value="")
t1 = [row("All detected genes", "RPKM > 0 in RNA-seq and Ribo-seq (≥1 tissue)", det),
      row("Canonical high-TE", "log2TE > 1.5 and log2 RNA ≥ 0 in every high-TE tissue", canon),
      row("Translationally primed", "log2 RNA < 0 and log2TE > 1.5 in the same tissue (≥1 tissue)", m.primed_any),
      row("   nuclear-encoded", "", m.primed_any & (m.organelle == "nuclear")),
      row("   primed in both tissues", "", m.primed_both),
      row("   primed with ≥10 Ribo-seq reads", "in a tissue where primed", m.high_confidence_primed)]
for t in ("root", "shoot"):
    base = m[m[f"ribo_detected_{t}"] & (m[f"log2_rna_{t}"] < 0) & (m[f"ribo_reads_{t}"] >= 10)]
    a, b = base[base[f"primed_{t}"]], base[~base[f"primed_{t}"]]
    f = fisher_exact([[a.protein_any.sum(), len(a) - a.protein_any.sum()], [b.protein_any.sum(), len(b) - b.protein_any.sum()]])
    t1.append(dict(gene_set=f"RNA-low genes, {t}s, ≥10 Ribo-seq reads: primed", criterion="log2TE > 1.5", n=len(a),
                   with_protein=int(a.protein_any.sum()), pct=round(100 * a.protein_any.mean(), 1),
                   odds_ratio=round(f.statistic, 2), p_value=f"{f.pvalue:.1e}"))
    t1.append(dict(gene_set=f"RNA-low genes, {t}s, ≥10 Ribo-seq reads: not primed", criterion="log2TE ≤ 1.5", n=len(b),
                   with_protein=int(b.protein_any.sum()), pct=round(100 * b.protein_any.mean(), 1), odds_ratio="(reference)", p_value=""))
t1 = pd.DataFrame(t1)
t1.to_csv(RESULTS / "Table1_protein_evidence.tsv", sep="\t", index=False)

# ---- KEGG ath00999 test ------------------------------------------------
kegg = set(open(RAW / "KEGG_ath00999_genes.txt").read().split())
nuc_det = m[(m.organelle == "nuclear") & (m.rna_detected_root | m.rna_detected_shoot)]
N, K = len(nuc_det), int(nuc_det.agi.isin(kegg).sum())
n, k = int(nuc_det.primed_any.sum()), int((nuc_det.primed_any & nuc_det.agi.isin(kegg)).sum())
p_hyper = hypergeom.sf(k - 1, N, K, n)
kegg_genes = tidy(m[m.primed_any & m.agi.isin(kegg)])
kegg_summary = pd.DataFrame({"Quantity": ["Background: nuclear genes detected by RNA-seq", "KEGG ath00999 genes in background",
                                          "Nuclear primed genes", "Nuclear primed genes in ath00999", "Expected", "Fold enrichment",
                                          "Hypergeometric P (one-sided)"],
                             "Value": [N, K, n, k, round(n * K / N, 2), round(k / (n * K / N), 2), f"{p_hyper:.2e}"]})

# ---- protein vs read support ----------------------------------------------
prot_reads = []
for t in ("root", "shoot"):
    base = m[m[f"ribo_detected_{t}"] & (m[f"log2_rna_{t}"] < 0)]
    for thr in (1, 5, 10, 20):
        a = base[base[f"primed_{t}"] & (base[f"ribo_reads_{t}"] >= thr)]
        b = base[~base[f"primed_{t}"] & (base[f"ribo_reads_{t}"] >= thr)]
        f = fisher_exact([[a.protein_any.sum(), len(a) - a.protein_any.sum()], [b.protein_any.sum(), len(b) - b.protein_any.sum()]])
        prot_reads.append({"Tissue": t, "Minimum Ribo-seq reads": thr, "Primed, n": len(a), "Primed with protein, n": int(a.protein_any.sum()),
                           "Primed with protein, %": round(100 * a.protein_any.mean(), 1), "RNA-low not primed, n": len(b),
                           "RNA-low not primed with protein, n": int(b.protein_any.sum()),
                           "RNA-low not primed with protein, %": round(100 * b.protein_any.mean(), 1),
                           "Odds ratio": round(f.statistic, 2), "Fisher P (two-sided)": f"{f.pvalue:.2e}"})
prot_reads = pd.DataFrame(prot_reads)

# ---- iron network --------------------------------------------------------
_fig2b = SI.parent / "figures" / "Fig2b_iron_genes.tsv"
iron = pd.read_csv(_fig2b, sep="\t") if _fig2b.exists() else None

# ---- Kozak matrices (Wen-Dar Lin) ---------------------------------------
mx = pd.read_excel(TE_FILE, sheet_name="ATG_region_matrixes", header=None)
kz = []
for i, v in mx[0].items():
    if isinstance(v, str) and "TE" in v:
        sub = mx.iloc[i + 1:i + 5].set_index(0).astype(float)
        sub.columns = ["-3", "-2", "-1", "A (+1)", "T (+2)", "G (+3)", "+4", "+5", "+6"]
        sub.index.name = "Nucleotide"; sub = sub.reset_index(); sub.insert(0, "TE set", v)
        kz.append(sub.round(3))
kz = pd.concat(kz)

# ---- g:Profiler enrichment (script 05) -------------------------------------
GPR = ROOT / "data" / "gprofiler" / "results"
GP_SETS = [("canonical_highTE_nuclear_154", "Canonical high-TE, nuclear (n = 154)", "genome"),
           ("primed_nuclear_701", "Translationally primed, nuclear (n = 701)", "genome"),
           ("primed_nuclear_701_bgDetected", "Translationally primed, nuclear (n = 701)", "nuclear genes detected by RNA-seq (n = 24,755)"),
           ("primed_both_nuclear_41", "Primed in both tissues, nuclear (n = 41)", "genome"),
           ("highTE_both_nuclear_93", "High-TE in both tissues, nuclear (n = 93)", "genome"),
           ("highTE_any_nuclear_855", "High-TE in >=1 tissue, nuclear (n = 855)", "genome")]
gp_sum, gp_terms = [], []
for key, label, bgd in GP_SETS:
    a = pd.read_csv(GPR / f"gprofiler_{key}_all.csv")
    sig = a[a.significant]
    # annotated query genes per source (reported by g:Profiler only for sources with >= 1 returned term;
    # not meaningful with a custom background, where query_size equals the query length)
    ann = a.groupby("source").query_size.first().to_dict() if bgd == "genome" else {}
    gp_sum.append({"Gene set": label, "Background": bgd, "Significant terms (FDR < 0.05)": len(sig),
                   "GO:BP": int((sig.source == "GO:BP").sum()), "GO:MF": int((sig.source == "GO:MF").sum()),
                   "GO:CC": int((sig.source == "GO:CC").sum()), "KEGG": int((sig.source == "KEGG").sum()),
                   "Query genes with GO:BP annotation": ann.get("GO:BP", ""), "Query genes with KEGG annotation": ann.get("KEGG", "")})
    t = sig[["source", "native", "name", "p_value", "term_size", "query_size", "intersection_size", "intersections"]].copy()
    t.insert(0, "Background", bgd); t.insert(0, "Gene set", label)
    t["p_value"] = t.p_value.map(lambda v: f"{v:.2e}")
    gp_terms.append(t)
gp_sum = pd.DataFrame(gp_sum)
gp_terms = pd.concat(gp_terms).rename(columns={"source": "Source", "native": "Term ID", "name": "Term", "p_value": "Adjusted P (BH)",
                                               "term_size": "Term size", "query_size": "Annotated query genes",
                                               "intersection_size": "Genes in term", "intersections": "Genes (AGI)"})
gp_info = (GPR / "run_info.txt").read_text().splitlines()[1]

# ---- 5'UTR motif analysis (script 04) + MEME low-TE motifs (original discovery) --
mot = pd.read_csv(RESULTS / "motif_enrichment.tsv", sep="\t", comment="#")
mot = mot.rename(columns={"set": "Gene set", "n": "Sequences (n)", "with_site": "With >=1 PWM site (n)", "pct": "With site (%)",
                          "regex_exact": "Exact CTATA[AT]A[AT]A match (n)", "regex_pct": "Exact match (%)",
                          "median_5utr_len": "Median annotated 5'UTR (nt)",
                          "OR_vs_low": "OR vs low-TE", "P_vs_low": "Fisher P vs low-TE", "OR_vs_lowmid": "OR vs low/mid-TE",
                          "P_vs_lowmid": "Fisher P vs low/mid-TE",
                          "5'UTR": "Sites in annotated 5'UTR", "upstream of annotated TSS": "Sites upstream of annotated TSS", "CDS": "Sites in CDS"})
motp = pd.read_csv(RESULTS / "motif_primed_upstream150.tsv", sep="\t").rename(columns={
    "set": "Gene set", "n": "Sequences (n)", "with_site": "With >=1 PWM site (n)", "pct": "With site (%)", "regex": "Exact match (n)",
    "regex_pct": "Exact match (%)", "sites_in_5utr": "Sites in annotated 5'UTR", "sites_upstream_TSS": "Sites upstream of annotated TSS",
    "pct_sites_upstream_TSS": "Sites upstream of TSS (%)", "median_5utr_len": "Median annotated 5'UTR (nt)",
    "pct_no_5utr": "No annotated 5'UTR (%)", "OR_vs_lowmid": "OR vs low/mid-TE", "P_vs_lowmid": "Fisher P vs low/mid-TE",
    "OR_vs_low": "OR vs low-TE", "P_vs_low": "Fisher P vs low-TE"})
for d in (mot, motp):
    for c in d.columns:
        if c.startswith("Fisher P"):
            d[c] = d[c].map(lambda v: "" if pd.isna(v) else f"{v:.2e}")
        elif d[c].dtype == float:
            d[c] = d[c].round(2)
msites = pd.read_csv(RESULTS / "motif_sites.tsv", sep="\t").rename(columns={
    "set": "Gene set", "gene_model": "Gene model", "start": "Site start (nt from A of AUG)", "centre": "Site centre (nt from A of AUG)",
    "utr5_len": "Annotated 5'UTR length (nt)", "location": "Location"})
meme = pd.read_excel(MOTIF_FILE, sheet_name="Motif_discovery_lowTE")

# ---- SEA enrichment (MEME Suite 5.4.1; W.-D. Lin) ----------------------------------
SEA = ROOT / "data" / "meme_sea"
SEA_RUNS = [  # (folder, target set, background set)
    ("sea_discovery/generalMid_TEhigh.utr5_SEA_1", "High-TE (log2TE > 1.5 in both tissues; discovery set), -150/+50", "Low-TE, -150/+50"),
    ("sea_discovery/generalMid_TEhigh.utr5_MidLow_SEA_1", "High-TE (log2TE > 1.5 in both tissues; discovery set), -150/+50", "Low/mid-TE (all), -150/+50"),
]


def read_sea(folder):
    t = pd.read_csv(SEA / folder / "sea.tsv", sep="\t", comment="#")
    return t.dropna(subset=["ID"])


sea_rows, sea_full = [], []
for folder, target, bgset in SEA_RUNS:
    t = read_sea(folder)
    r = t[t.CONSENSUS == "CTATAWAWA"]
    if len(r):
        r = r.iloc[0]
        sea_rows.append({"Target sequences": target, "Background sequences": bgset,
                         "Target with motif (n)": int(r.TP), "Target with motif (%)": r["TP%"],
                         "Background with motif (n)": int(r.FP), "Background with motif (%)": r["FP%"],
                         "Enrichment ratio": r.ENR_RATIO, "Score threshold": r.SCORE_THR,
                         "P value": f"{r.PVALUE:.2e}", "E value": f"{r.EVALUE:.2e}", "SEA output": folder})
    if folder == "sea_discovery/generalMid_TEhigh.utr5_SEA_1":
        sea_full = t.drop(columns=["DB"]).assign(**{"SEA output": folder})
sea_tab = pd.DataFrame(sea_rows)

# ---- assemble workbook --------------------------------------------------
readme = [
    ("Supporting Information Tables S1-S9", ""),
    ("", ""),
    ("Data sources", "RNA-seq and Ribo-seq: Vélez-Bermúdez et al. (2025) Scientific Data 12: 504 (processed per-tissue RPKM values, file TranslationEfficiency2.xlsx, W.-D. Lin). "
                     "Proteomes: Jain & Schmidt (2023) Molecular & Cellular Proteomics, supplementary dataset (root and shoot proteome sheets). Annotation: Araport11."),
    ("TE", "Translational efficiency = Ribo-seq RPKM / RNA-seq RPKM per tissue. Requires >=1 RNA-seq read (not defined when RNA-seq RPKM = 0); no higher read-count threshold was applied. "
           "Genes with RNA-seq reads but no Ribo-seq reads have TE = 0 and are classified low-TE (flag column 'No Ribo-seq reads (TE = 0)')."),
    ("TE classes", "high-TE: log2TE > 1.5; mid-TE: -1.5 <= log2TE <= 1.5; low-TE: log2TE < -1.5 (per tissue)."),
    ("Translationally primed", "log2 RNA-seq RPKM < 0 (RPKM < 1) and log2TE > 1.5 in the same tissue (therefore >=1 RNA-seq and >=1 Ribo-seq read). "
                               "'Primed with >=10 Ribo-seq reads' = supported by >=10 Ribo-seq reads in a tissue in which the gene is primed."),
    ("Canonical high-TE", "log2TE > 1.5 with log2 RNA-seq RPKM >= 0 in every tissue in which the gene is high-TE."),
    ("Protein evidence", "Protein detected with High or Medium FDR confidence in the root or shoot proteome (any pH condition), matched by AGI locus."),
    ("", ""),
    ("Numbering", "Tables are numbered in the order in which they are first cited in the main text."),
    ("Table S1", "All 27,655 gene models: read counts, RPKM, log2 values, TE, TE classes, priming status and protein evidence."),
    ("Table S2", "High-TE genes (log2TE > 1.5 in >=1 tissue; n = 949) with priming status."),
    ("Table S3", f"GO and KEGG enrichment (g:Profiler g:GOSt, {gp_info}; GO:BP, GO:MF, GO:CC and KEGG; electronic GO annotations excluded; "
                 "Benjamini-Hochberg FDR < 0.05; background = annotated genome, and, for the primed set, also all nuclear genes detected by RNA-seq). "
                 "(A) summary per gene set; (B) all significant terms with the genes in each term. Sheet 'Table S3 (C) ath00999': primed genes in KEGG pathway ath00999 "
                 "and an independent one-sided hypergeometric test against nuclear genes detected by RNA-seq (KEGG REST gene list)."),
    ("Table S4", "Translationally primed genes (n = 757; 701 nuclear-encoded)."),
    ("Table S5", "Protein-level evidence for RNA-low genes as a function of Ribo-seq read support (source data for Fig. 1e,f), and summary of protein evidence by gene set."),
    ("Table S6", "Nuclear-encoded genes primed in both roots and shoots (n = 41), ranked by mean log2TE (source data for Fig. 2a)."),
    ("Table S7", "Iron-acquisition network genes (source data for Fig. 2b)."),
    ("Table S8", "Nucleotide frequencies around the AUG start codon by TE set (W.-D. Lin; source data for Fig. 2c). "
                 "Sets restricted to genes with log2 RNA-seq RPKM >= 2 in both tissues; thresholds > 0.5 and > 1.0 are nested and include the high-TE set."),
    ("Table S9", "CTATA[A/T]A[A/T]A-like motif (MEME motif 6; MEME v.4.12.0, discovered in the -150/-51 region of the 27 high-TE sequences) - source data for Fig. S1. "
                 "(A) SEA (MEME Suite v.5.4.1; W.-D. Lin) of the MEME motifs in the discovery set (high-TE vs low-TE and vs low/mid-TE sequences), used to select motif 6. "
                 "(B) PWM scan (site P < 1e-4, sense strand) of the -150/+50 window in the discovery set, two independent sets not used for discovery "
                 "(1.0 < log2TE <= 1.5 and 0.5 < log2TE <= 1.5 in both tissues), the remaining low/mid-TE and the low-TE genes; two-sided Fisher's exact tests against "
                 "the remaining low/mid-TE and the low-TE genes (real sequences); location of sites relative to the annotated 5'UTR. "
                 "(C) Same scan of the -150/-1 window of the 701 nuclear primed genes; (E) motifs discovered by MEME in low-TE sequences. "
                 "Sheet 'Table S9 (D) motif sites': position and location of every site."),
    ("", ""),
    ("Code", "All tables and figures are generated by the scripts in the accompanying GitHub repository (scripts/01-05)."),
]
out = SI / "Supporting_Information_Tables_S1-S9.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    pd.DataFrame(readme, columns=["Item", "Description"]).to_excel(xw, sheet_name="README", index=False)
    tidy(m).to_excel(xw, sheet_name="Table S1 all genes", index=False)
    tidy(m[m.high_te_any].sort_values("mean_log2_te", ascending=False)).to_excel(xw, sheet_name="Table S2 high-TE", index=False)
    gp_sum.to_excel(xw, sheet_name="Table S3 GO-KEGG enrichment", index=False)
    gp_terms.to_excel(xw, sheet_name="Table S3 GO-KEGG enrichment", index=False, startrow=len(gp_sum) + 3)
    kegg_summary.to_excel(xw, sheet_name="Table S3 (C) ath00999", index=False)
    kegg_genes.to_excel(xw, sheet_name="Table S3 (C) ath00999", index=False, startrow=len(kegg_summary) + 3)
    tidy(m[m.primed_any].sort_values("mean_log2_te", ascending=False)).to_excel(xw, sheet_name="Table S4 primed", index=False)
    prot_reads.to_excel(xw, sheet_name="Table S5 protein vs reads", index=False)
    t1.rename(columns={"gene_set": "Gene set", "criterion": "Criterion", "n": "Genes (n)", "with_protein": "With protein evidence (n)",
                       "pct": "With protein evidence (%)", "odds_ratio": "Odds ratio", "p_value": "Fisher P"}).to_excel(
        xw, sheet_name="Table S5 protein vs reads", index=False, startrow=len(prot_reads) + 3)
    s6 = m[m.primed_both & (m.organelle == "nuclear")].sort_values("mean_log2_te", ascending=False)
    tidy(s6, extra=("mean_log2_te",)).rename(columns={"mean_log2_te": "Mean log2 TE"}).to_excel(xw, sheet_name="Table S6 shared nuclear", index=False)
    if iron is not None:
        iron.rename(columns={"gene": "Gene", "agi": "AGI locus", "group": "Group", "tissue": "Tissue", "log2_rna": "log2 RNA-seq RPKM",
                             "log2_te": "log2 TE", "primed": "Translationally primed", "ribo_reads": "Ribo-seq reads",
                             "rna_reads": "RNA-seq reads", "protein": "Protein detected (root or shoot)"}).to_excel(
            xw, sheet_name="Table S7 iron network", index=False)
    kz.to_excel(xw, sheet_name="Table S8 start-codon context", index=False)
    sea_tab.to_excel(xw, sheet_name="Table S9 (A) SEA", index=False)
    if len(sea_full):
        sea_full.to_excel(xw, sheet_name="Table S9 (A) SEA", index=False, startrow=len(sea_tab) + 3)
    mot.to_excel(xw, sheet_name="Table S9 motif", index=False)
    motp.to_excel(xw, sheet_name="Table S9 motif", index=False, startrow=len(mot) + 3)
    meme.to_excel(xw, sheet_name="Table S9 motif", index=False, startrow=len(mot) + len(motp) + 6)
    msites.to_excel(xw, sheet_name="Table S9 (D) motif sites", index=False)

wb = load_workbook(out)
hdr = PatternFill("solid", fgColor="E8E7E3")
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w * 0.9), 60)
ws = wb["README"]; ws.column_dimensions["B"].width = 140
for r in ws.iter_rows(min_row=2):
    r[1].alignment = Alignment(wrap_text=True, vertical="top")
    if r[0].value and str(r[0].value).startswith("Supporting"):
        r[0].font = Font(bold=True, size=13)
for name, title_row in (("Table S3 (C) ath00999", len(kegg_summary) + 4), ("Table S5 protein vs reads", len(prot_reads) + 4),
                        ("Table S3 GO-KEGG enrichment", len(gp_sum) + 4), ("Table S9 motif", len(mot) + 4),
                        ("Table S9 motif", len(mot) + len(motp) + 7), ("Table S9 (A) SEA", len(sea_tab) + 4)):
    for c in wb[name][title_row]:
        c.font = Font(bold=True); c.fill = hdr
# section labels
for name, row, text in (("Table S3 GO-KEGG enrichment", len(gp_sum) + 3, "(B) Significant terms (FDR < 0.05)"),
                        ("Table S9 motif", len(mot) + 3, "(C) Translationally primed genes, -150/-1 window, compared with the same window of the sets in (B)"),
                        ("Table S9 motif", len(mot) + len(motp) + 6, "(E) Motifs discovered by MEME in low-TE sequences (original analysis)"),
                        ("Table S9 (A) SEA", len(sea_tab) + 3, "All MEME motifs of the discovery run (high-TE, -150/-51) tested with SEA against low-TE sequences")):
    wb[name].cell(row=row, column=1, value=text).font = Font(bold=True, italic=True)
ws = wb["Table S3 GO-KEGG enrichment"]; ws.column_dimensions["J"].width = 60
# ---- table headings (one title row above each sheet) --------------------------
TITLES = {
    "Table S1 all genes": "Table S1 All gene models: read counts, RPKM, translational efficiency (TE), TE classes, priming status and protein evidence.",
    "Table S2 high-TE": "Table S2 High-TE genes (log₂TE > 1.5 in at least one tissue; n = 949) with priming status.",
    "Table S3 GO-KEGG enrichment": "Table S3 Gene Ontology and KEGG enrichment analyses (g:Profiler). (A) Summary per gene set; (B) significant terms (FDR < 0.05) with the genes in each term.",
    "Table S3 (C) ath00999": "Table S3 (continued) (C) Translationally primed genes in KEGG pathway ath00999 (Biosynthesis of various plant secondary metabolites) and one-sided hypergeometric test.",
    "Table S4 primed": "Table S4 Translationally primed genes (log₂ RNA-seq RPKM < 0 and log₂TE > 1.5 in the same tissue; n = 757, of which 701 are nuclear-encoded).",
    "Table S5 protein vs reads": "Table S5 Protein-level evidence for RNA-low genes as a function of Ribo-seq read support (top; source data for Fig. 1e,f) and summary of protein evidence by gene set (bottom).",
    "Table S6 shared nuclear": "Table S6 Nuclear-encoded genes primed in both roots and shoots (n = 41), ranked by mean log₂TE (source data for Fig. 2a).",
    "Table S7 iron network": "Table S7 Iron-acquisition network genes: TE, transcript abundance, read support and protein evidence in roots and shoots (source data for Fig. 2b).",
    "Table S8 start-codon context": "Table S8 Nucleotide frequencies around the AUG start codon by TE class (source data for Fig. 2c).",
    "Table S9 (A) SEA": "Table S9 CTATA[A/T]A[A/T]A-like motif (source data for Fig. S1). (A) Enrichment of the MEME motifs in high-TE sequences (SEA, MEME Suite v.5.4.1).",
    "Table S9 motif": "Table S9 (continued) (B) Motif prevalence (PWM scan) and Fisher's exact tests, −150 to +50 nt; (C) translationally primed genes, −150 to −1 nt; (E) motifs discovered by MEME in low-TE sequences.",
    "Table S9 (D) motif sites": "Table S9 (continued) (D) Position and location of every motif site.",
}
for name, title in TITLES.items():
    ws = wb[name]
    ws.insert_rows(1)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=12); c.alignment = Alignment(wrap_text=False, vertical="center")
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A3"
wb.save(out)
print("written", out)
print(t1.to_string())
print(kegg_summary.to_string())
print(gp_sum.to_string())
